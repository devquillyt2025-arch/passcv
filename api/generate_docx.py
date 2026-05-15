"""
Vercel Python serverless function — DOCX generator.
Accepts: POST with JSON { "resume": RewrittenResume, "jobTitle": string }
Returns: DOCX binary (application/vnd.openxmlformats-officedocument.wordprocessingml.document)
"""
from http.server import BaseHTTPRequestHandler
import json
import io
import base64


def build_docx(resume: dict, job_title: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # Page margins — generous but not wasteful
    for section in doc.sections:
        section.top_margin    = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin   = Cm(2.0)
        section.right_margin  = Cm(2.0)

    styles = doc.styles

    # Default body font
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)

    def add_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(2)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)  # brand blue
        # Bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '1E3A8A')
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def add_body(text: str, bold: bool = False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(1)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(10.5)
        return p

    def add_bullet(text: str):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(1)
        p.paragraph_format.left_indent  = Cm(0.5)
        run = p.add_run(text)
        run.font.size = Pt(10.5)
        return p

    contact = resume.get('contact', {})

    # ── Name ──────────────────────────────────────────────────────────────
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(contact.get('name', ''))
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # ── Contact line ──────────────────────────────────────────────────────
    contact_parts = [
        contact.get('email', ''),
        contact.get('phone', ''),
        contact.get('location', ''),
        contact.get('linkedin', ''),
    ]
    if resume.get('noticePeriod'):
        contact_parts.append(f"Notice Period: {resume['noticePeriod']}")
    contact_line = '  |  '.join(p for p in contact_parts if p)
    ct_para = doc.add_paragraph()
    ct_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ct_run = ct_para.add_run(contact_line)
    ct_run.font.size = Pt(9.5)
    ct_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ── Summary ───────────────────────────────────────────────────────────
    if resume.get('summary'):
        add_heading('Summary')
        add_body(resume['summary'])

    # ── Skills ────────────────────────────────────────────────────────────
    if resume.get('skills'):
        add_heading('Skills')
        add_body(', '.join(resume['skills']))

    # ── Work Experience ───────────────────────────────────────────────────
    if resume.get('experience'):
        add_heading('Work Experience')
        for job in resume['experience']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(1)
            title_run = p.add_run(f"{job.get('title', '')}  —  {job.get('company', '')}")
            title_run.bold = True
            title_run.font.size = Pt(10.5)

            date_str = f"{job.get('startDate', '')} – {job.get('endDate', '')}"
            date_run = p.add_run(f"  ({date_str})")
            date_run.font.size = Pt(10)
            date_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            for bullet in job.get('bullets', []):
                add_bullet(bullet)

    # ── Education ─────────────────────────────────────────────────────────
    if resume.get('education'):
        add_heading('Education')
        for edu in resume['education']:
            deg_line = f"{edu.get('degree', '')} in {edu.get('field', '')}".strip(' in ')
            inst_line = f"{edu.get('institution', '')}  ({edu.get('year', '')})"
            cgpa = edu.get('cgpa', '')
            if cgpa:
                inst_line += f"  —  CGPA/Marks: {cgpa}"
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(1)
            p.add_run(deg_line).bold = True
            doc.add_paragraph(inst_line).paragraph_format.space_after = Pt(1)

    # ── Certifications ────────────────────────────────────────────────────
    if resume.get('certifications'):
        add_heading('Certifications')
        for cert in resume['certifications']:
            add_bullet(cert)

    # ── Declaration ───────────────────────────────────────────────────────
    doc.add_paragraph()
    decl = doc.add_paragraph()
    decl_run = decl.add_run(
        "I hereby declare that all information provided above is true and correct to the best of my knowledge."
    )
    decl_run.font.size = Pt(9.5)
    decl_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Vercel handler ──────────────────────────────────────────────────────────

class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            data = json.loads(body)

            resume   = data['resume']
            job_title = data.get('jobTitle', 'Resume')
            name     = resume.get('contact', {}).get('name', 'Resume')
            print('DOCX generator request', {
                'jobTitle': job_title,
                'name': name,
                'resume_fields': list(resume.keys()) if isinstance(resume, dict) else None,
            })

            docx_bytes = build_docx(resume, job_title)

            safe_name     = re.sub(r'[^a-zA-Z0-9_-]', '_', name)[:30]
            safe_title    = re.sub(r'[^a-zA-Z0-9_-]', '_', job_title)[:30]
            filename      = f"{safe_name}_{safe_title}_TailorCV.docx"

            self.send_response(200)
            self.send_header('Content-Type',
                             'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            self.send_header('Content-Disposition', f'attachment; filename="{filename}"')
            self.send_header('Content-Length', str(len(docx_bytes)))
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(docx_bytes)

        except Exception as e:
            err = json.dumps({'error': str(e)}).encode()
            self.send_response(500)
            self.send_header('Content-Type', 'application/json')
            self.send_header('Content-Length', str(len(err)))
            self.end_headers()
            self.wfile.write(err)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def log_message(self, format, *args):
        pass


import re  # noqa: E402 — import at top in prod; here for inline Vercel bundling
