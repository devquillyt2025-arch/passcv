from __future__ import annotations

import io
import json
import os
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import ValidationError

from .schemas import ResumeData


ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


app = FastAPI(
    title="Resume Parser Service",
    version="1.0.0",
    description="Parses PDF and DOCX resumes into a strict frontend-ready JSON schema.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["POST", "GET"],
    allow_headers=["*"]
)


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.post("/api/parse-resume", response_model=ResumeData)
async def parse_resume(file: UploadFile = File(...)) -> ResumeData:
    if file.content_type not in ALLOWED_MIME_TYPES and not file.filename.lower().endswith((".pdf", ".docx")):
        raise HTTPException(status_code=400, detail="Unsupported file type. Only PDF and DOCX are accepted.")

    try:
        contents = await file.read()
        raw_text = extract_text(contents, file.content_type, file.filename)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Unable to read the uploaded file.") from exc

    if not raw_text.strip():
        raise HTTPException(status_code=400, detail="The uploaded resume contains no readable text.")

    try:
        parsed_json = llm_parse_resume(raw_text)
        parsed_resume = ResumeData.model_validate(parsed_json)
        return parsed_resume
    except ValidationError as exc:
        raise HTTPException(status_code=400, detail=f"Parsed output did not match the resume schema: {exc}") from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Failed to parse resume with the AI model.") from exc


def extract_text(file_bytes: bytes, content_type: str, filename: str) -> str:
    if content_type == "application/pdf" or filename.lower().endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)

    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.lower().endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber is required to parse PDF resumes.") from exc

    full_text = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=2, y_tolerance=2)
            if text:
                full_text.append(text)

    return "\n".join(full_text).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to parse DOCX resumes.") from exc

    document = Document(io.BytesIO(file_bytes))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    tables: list[str] = []
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                tables.append(row_text)

    return "\n".join(paragraphs + tables).strip()


def llm_parse_resume(raw_text: str) -> dict[str, Any]:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OpenAI API key is not configured. Set OPENAI_API_KEY to enable AI parsing.")

    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("OpenAI Python SDK is required for AI parsing.") from exc

    client = OpenAI(api_key=api_key)
    model = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")

    schema = ResumeData.model_json_schema()
    system_prompt = (
        "Extract the resume into a strict JSON object matching the frontend schema exactly. "
        "Do not include markdown, explanation, or any extra fields. "
        "Return only valid JSON. "
        "If a field is missing, use an empty string or an empty array. "
        "Use the schema and example object shape exactly as provided."
    )

    user_prompt = (
        "Here is the raw resume text. Parse all contact, summary, experience, education, skills, "
        "projects, certifications, and languages into the JSON schema. "
        "For each experience item, place bullets into the `description` field as a single string with newline-separated lines. "
        "If a value is not present on the resume, return an empty string or an empty list. "
        "Raw resume text:\n\n" + raw_text
    )

    response = client.responses.create(
        model=model,
        input=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object", "json_schema": schema},
    )

    parsed = None
    if hasattr(response, "output_parsed") and response.output_parsed is not None:
        parsed = response.output_parsed
    else:
        text_output = None
        if hasattr(response, "output"):
            output = response.output
            if isinstance(output, list) and output and isinstance(output[0], dict):
                text_output = output[0].get("content", [{}])[0].get("text")
        if not text_output:
            raise ValueError("AI returned no valid JSON payload.")
        parsed = json.loads(text_output)

    if not isinstance(parsed, dict):
        raise ValueError("AI did not return a JSON object.")

    return parsed
