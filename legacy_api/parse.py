"""
Vercel Python serverless function — resume parser.
Accepts: POST with JSON { "content": "<base64>", "filename": "resume.pdf" }
Returns: ParsedResume JSON
"""
from http.server import BaseHTTPRequestHandler
import json
import base64
import io
import re


# ── helpers ────────────────────────────────────────────────────────────────

DATE_PAT = re.compile(
    r'(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|'
    r'Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'
    r'\s*(\d{4})\s*[-–—to]+\s*'
    r'(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|'
    r'Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?'
    r'|Present|Current|Till date|Till Date)?\s*(\d{4})?',
    re.I,
)

SECTION_MAP = {
    'summary':          ['summary', 'professional summary', 'profile', 'about me',
                         'objective', 'career objective', 'professional profile'],
    'experience':       ['experience', 'work experience', 'employment history',
                         'work history', 'professional experience', 'career history',
                         'internship', 'internships'],
    'education':        ['education', 'academic background', 'academic qualifications',
                         'educational background', 'qualification'],
    'skills':           ['skills', 'technical skills', 'key skills', 'core competencies',
                         'competencies', 'expertise', 'technologies'],
    'certifications':   ['certifications', 'certificates', 'credentials', 'training',
                         'courses', 'achievements', 'awards'],
}


def parse_text(full_text: str, has_multi_col: bool, has_tables: bool, has_images: bool) -> dict:
    lines = [l.rstrip() for l in full_text.split('\n')]
    non_empty = [l for l in lines if l.strip()]

    result = {
        'contact': {'name': '', 'email': '', 'phone': '', 'location': '', 'linkedin': ''},
        'summary': '',
        'experience': [],
        'education': [],
        'skills': [],
        'certifications': [],
        'noticePeriod': None,
        'ctc': None,
        'hasMultiColumn': has_multi_col,
        'hasTables': has_tables,
        'hasImages': has_images,
        'rawText': full_text[:5000],
    }

    # Contact extraction
    email_m = re.search(r'\b[\w.%+-]+@[\w.-]+\.[a-z]{2,}\b', full_text, re.I)
    if email_m:
        result['contact']['email'] = email_m.group()

    phone_m = re.search(r'(?:\+91[-\s]?)?[6-9]\d{9}|\(\d{3,4}\)\s?\d{6,8}|\d{5}[\s-]\d{5}', full_text)
    if phone_m:
        result['contact']['phone'] = phone_m.group()

    li_m = re.search(r'linkedin\.com/in/([\w-]+)', full_text, re.I)
    if li_m:
        result['contact']['linkedin'] = f"linkedin.com/in/{li_m.group(1)}"

    # Name: first non-empty line that has no @, digits, or punctuation-only chars
    for line in non_empty[:5]:
        stripped = line.strip()
        if stripped and not re.search(r'[@\d|•\-_=]', stripped) and len(stripped.split()) <= 6 and len(stripped) > 3:
            result['contact']['name'] = stripped
            break

    # Location (Indian cities)
    city_m = re.search(
        r'\b(Mumbai|Delhi|Bangalore|Bengaluru|Hyderabad|Chennai|Pune|Kolkata|'
        r'Ahmedabad|Noida|Gurgaon|Gurugram|Jaipur|Bhopal|Indore|Lucknow|'
        r'Chandigarh|Kochi|Coimbatore|Surat|Nagpur|Vadodara|Patna|Bhubaneswar)\b',
        full_text, re.I
    )
    if city_m:
        result['contact']['location'] = city_m.group(1)

    # Notice period & CTC
    notice_m = re.search(r'notice\s*period\s*[:\-]?\s*(.+)', full_text, re.I)
    if notice_m:
        result['noticePeriod'] = notice_m.group(1).strip()[:60]

    ctc_m = re.search(r'(?:current\s+)?ctc\s*[:\-]?\s*(.+)', full_text, re.I)
    if ctc_m:
        result['ctc'] = ctc_m.group(1).strip()[:60]

    # Section detection
    current_section = None
    section_content: dict = {k: [] for k in SECTION_MAP}

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        lower = stripped.lower()
        found = None
        for sec, headers in SECTION_MAP.items():
            if any(lower == h or lower.startswith(h + ':') or lower.startswith(h + ' ') for h in headers):
                found = sec
                break

        if found:
            current_section = found
        elif current_section:
            section_content[current_section].append(stripped)

    # Summary
    result['summary'] = ' '.join(section_content['summary'][:4])

    # Skills
    skills_raw = ' '.join(section_content['skills'])
    skills = re.split(r'[,|•·\n/]', skills_raw)
    result['skills'] = [s.strip() for s in skills if 2 < len(s.strip()) < 50][:50]

    # Certifications
    result['certifications'] = [l for l in section_content['certifications'] if len(l) < 120][:10]

    # Experience
    result['experience'] = parse_experience(section_content['experience'])

    # Education
    result['education'] = parse_education(section_content['education'])

    return result


def parse_experience(lines: list) -> list:
    jobs = []
    current: dict | None = None

    for line in lines:
        date_m = DATE_PAT.search(line)
        is_bullet = line.startswith(('•', '-', '·', '◦', '▪', '*', '–'))
        is_short = len(line) < 120

        if date_m and is_short and not is_bullet:
            if current:
                jobs.append(current)
            current = {'company': '', 'title': '', 'startDate': '', 'endDate': '', 'bullets': []}

            # Parse dates
            g = date_m.groups()
            start_month = g[0] or ''
            start_year = g[1] or ''
            end_part = g[2] or ''
            end_year = g[3] or ''
            current['startDate'] = f"{start_month[:3]} {start_year}".strip()
            current['endDate'] = (f"{end_part[:3]} {end_year}".strip() if end_part else 'Present')

            # Title / company from non-date part of line
            pre = line[:date_m.start()].strip()
            pre = re.sub(r'\s+', ' ', pre).strip()
            if '|' in pre:
                parts = [p.strip() for p in pre.split('|')]
                current['title'] = parts[0]
                current['company'] = parts[1] if len(parts) > 1 else ''
            elif ',' in pre:
                parts = [p.strip() for p in pre.split(',', 1)]
                current['title'] = parts[0]
                current['company'] = parts[1] if len(parts) > 1 else ''
            else:
                current['title'] = pre

        elif current is not None:
            if is_bullet:
                bullet = re.sub(r'^[•\-·◦▪\*–]\s*', '', line).strip()
                if bullet:
                    current['bullets'].append(bullet)
            elif not current['company'] and len(line) < 80 and not line[0].islower():
                current['company'] = line.strip()
            elif len(line) > 20:
                current['bullets'].append(line.strip())

    if current:
        jobs.append(current)

    return jobs[:10]


def parse_education(lines: list) -> list:
    edu = []
    current: dict | None = None

    year_pat = re.compile(r'\b(19|20)\d{2}\b')

    for line in lines:
        year_m = year_pat.search(line)
        if year_m and len(line) < 120:
            if current:
                edu.append(current)
            current = {'institution': '', 'degree': '', 'field': '', 'year': year_m.group(), 'cgpa': ''}

            text = line[:year_m.start()].strip()
            # Look for CGPA
            cgpa_m = re.search(r'(?:cgpa|gpa|percentage|%)\s*[:\-]?\s*([\d.]+)', line, re.I)
            if cgpa_m:
                current['cgpa'] = cgpa_m.group(1)

            # Degree patterns
            deg_m = re.search(
                r'\b(B\.?Tech|M\.?Tech|BE|ME|BSc|MSc|MBA|BCA|MCA|B\.Com|M\.Com|BA|MA|PhD|Diploma|B\.E|M\.E)\b',
                text, re.I
            )
            if deg_m:
                current['degree'] = deg_m.group()
                current['field'] = text[deg_m.end():].strip().lstrip('.,- ')[:80]
            else:
                current['degree'] = text[:80]

        elif current is not None and not current['institution']:
            current['institution'] = line.strip()[:120]

    if current:
        edu.append(current)

    return edu[:5]


def parse_pdf(file_bytes: bytes) -> dict:
    import pdfplumber

    has_multi_col = False
    has_tables = False
    has_images = False
    full_text = ''

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            if page.images:
                has_images = True

            if page.find_tables():
                has_tables = True

            words = page.extract_words()
            if words and not has_multi_col:
                pw = page.width
                left_words  = [w for w in words if w['x0'] < pw * 0.45]
                right_words = [w for w in words if w['x0'] > pw * 0.55]
                if len(left_words) > 15 and len(right_words) > 15:
                    lmax = max(w['x1'] for w in left_words)
                    rmin = min(w['x0'] for w in right_words)
                    if rmin - lmax > pw * 0.04:
                        has_multi_col = True

            text = page.extract_text(x_tolerance=3, y_tolerance=3)
            if text:
                full_text += text + '\n'

    if not full_text.strip():
        return {'error': 'No text found. This may be a scanned/image-based PDF — please use a text-based PDF or DOCX.'}

    return parse_text(full_text, has_multi_col, has_tables, has_images)


def parse_docx(file_bytes: bytes) -> dict:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))

    has_tables = len(doc.tables) > 0
    has_images = any(
        'image' in (r.element.xml if hasattr(r.element, 'xml') else '')
        for p in doc.paragraphs for r in p.runs
    )

    # Check for multi-column sections
    has_multi_col = False
    for section in doc.sections:
        if section.start_type and hasattr(section, '_sectPr'):
            cols = section._sectPr.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols')
            if cols:
                num = cols[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '1')
                if int(num) > 1:
                    has_multi_col = True

    full_text = '\n'.join(p.text for p in doc.paragraphs)
    # Include table text
    for table in doc.tables:
        for row in table.rows:
            full_text += '\n' + ' | '.join(cell.text for cell in row.cells)

    if not full_text.strip():
        return {'error': 'Could not extract text from this DOCX file.'}

    return parse_text(full_text, has_multi_col, has_tables, has_images)


# ── Vercel handler ──────────────────────────────────────────────────────────

class handler(BaseHTTPRequestHandler):
    def do_POST(self):
        try:
            length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(length)
            data = json.loads(body)

            file_content = base64.b64decode(data['content'])
            filename = data.get('filename', 'resume.pdf').lower()

            if len(file_content) > 2 * 1024 * 1024:
                self._respond(400, {'error': 'File exceeds 2 MB limit. Please compress or use a smaller file.'})
                return

            if filename.endswith('.pdf'):
                result = parse_pdf(file_content)
            elif filename.endswith('.docx'):
                result = parse_docx(file_content)
            else:
                self._respond(400, {'error': 'Unsupported file type. Please upload a PDF or DOCX file.'})
                return

            self._respond(200, result)

        except Exception as e:
            self._respond(500, {'error': str(e)})

    def _respond(self, status: int, data: dict):
        body = json.dumps(data).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(body)))
        self.send_header('Access-Control-Allow-Origin', '*')
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def log_message(self, format, *args):
        pass
